import sys
from pathlib import Path

sys.path.append(
    str(Path(__file__).resolve().parent.parent)
)


from openpyxl import load_workbook

from app.database.database import SessionLocal
from app.models.province import Province
from app.models.city import City


PROVINCES_FILE = r"D:\استان ها.xlsx"
CITIES_FILE = r"D:\شهر های مربوط به هر استان.xlsx"


def import_provinces(db):
    workbook = load_workbook(PROVINCES_FILE)

    sheet = workbook.active

    provinces = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        name = row[0]

        if name:
            name = str(name).strip()

            province = Province(
                name=name
            )

            db.add(province)
            db.flush()

            provinces[name] = province.id

    db.commit()

    return provinces


def import_cities(db, provinces):
    workbook = load_workbook(CITIES_FILE)

    sheet = workbook.active

    cities_count = 0

    for row in sheet.iter_rows(min_row=2, values_only=True):

        province_name = row[0]
        city_name = row[1]

        if province_name and city_name:

            province_name = str(province_name).strip()
            city_name = str(city_name).strip()

            province_id = provinces.get(province_name)

            if province_id:

                city = City(
                    name=city_name,
                    province_id=province_id
                )

                db.add(city)
                cities_count += 1

    db.commit()

    return cities_count


def main():

    db = SessionLocal()

    try:

        print("Importing provinces...")

        provinces = import_provinces(db)

        print(
            f"{len(provinces)} provinces imported"
        )


        print("Importing cities...")

        cities = import_cities(
            db,
            provinces
        )

        print(
            f"{cities} cities imported"
        )


    finally:
        db.close()


if __name__ == "__main__":
    main()